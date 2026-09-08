from __future__ import annotations

import pandas as pd
from openpyxl import load_workbook

from src.core.export import export_workbook


def test_export_workbook_is_compact_and_paper_facing(tmp_path):
    output = tmp_path / "specimen_CrackVision.xlsx"
    frame_df = pd.DataFrame(
        [
            {
                "Frame": 184,
                "MTS_peak_force_N": 3521.7776,
                "MTS_peak_time_s": 922.169,
                "DIC_selected_time_s": 920.0,
                "frame_match_error_s": -2.169,
                "crack_count": 3,
                "Crack_width_mean_um": 53.333,
                "Crack_width_median_um": 40.0,
                "Crack_width_95_um": 94.0,
                "Crack_width_max_um": 100.0,
                "cod_status": "ok",
                "pixel_size_mm": 0.045,
                "dic_step_px": 3.0,
                "dic_point_spacing_mm": 0.135,
                "valid_fraction": 0.98,
                "principal_strain_threshold": 0.002,
                "candidate_points": 120,
                "skeleton_points": 80,
                "metadata_source": "synthetic",
            }
        ]
    )
    crack_df = pd.DataFrame(
        [
            {"Crack_ID": 1, "Length_mm": 8.2, "W_median_um": 20.0, "COD_samples": 100, "Fit_R2_median": 0.95},
            {"Crack_ID": 2, "Length_mm": 5.7, "W_median_um": 40.0, "COD_samples": 5, "Fit_R2_median": 0.92},
            {"Crack_ID": 3, "Length_mm": 9.4, "W_median_um": 100.0, "COD_samples": 3, "Fit_R2_median": 0.90},
        ]
    )

    export_workbook(output, frame_df, crack_df)
    wb = load_workbook(output)

    assert wb.sheetnames == ["01_结果汇总", "02_裂缝明细", "03_质量检查"]

    summary = wb["01_结果汇总"]
    assert "峰值拉应力状态裂缝宽度" in summary["A1"].value
    assert summary["C8"].value == frame_df.loc[0, "Crack_width_mean_um"]
    assert summary["C11"].value == "ok"
    assert len(summary._charts) == 1

    details = wb["02_裂缝明细"]
    assert [details.cell(1, col).value for col in range(1, 6)] == [
        "裂缝编号",
        "裂缝长度 (mm)",
        "代表宽度 (μm)",
        "COD 有效点数",
        "拟合 R² 中位数",
    ]
    assert details.max_row == 4
    assert details.cell(2, 3).value == 20.0
    assert details.cell(4, 3).value == 100.0

    qa = wb["03_质量检查"]
    assert qa["A1"].value == "检查项"
    assert qa["B1"].value == "数值"
